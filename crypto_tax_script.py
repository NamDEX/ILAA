import pandas as pd
import yfinance as yf
from decimal import Decimal, ROUND_HALF_UP
import datetime
import os
import sys
import openpyxl
from openpyxl.utils import get_column_letter

### USER CONFIGURATION ###
COINBASE_FILE_PATH = r"C:\Projects Python\Crypto DEC25\Input\coinbase.csv"
BINANCE_FILE_PATH = r"C:\Projects Python\Crypto DEC25\Input\binance.csv"
OUTPUT_FILE_NAME = r"C:\Projects Python\Crypto DEC25\Output\master_crypto_ledger.csv"

# Configure logging for clear console output
import logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

def fetch_fx_history():
    """
    Fetches the full history of GBPUSD=X from Yahoo Finance.
    Returns a Series with Date index and Close price.
    """
    logger.info("Fetching historical FX rates for GBPUSD=X...")
    try:
        # GBPUSD=X: Returns USD per 1 GBP.
        ticker = "GBPUSD=X"
        # Fetch MAX history to ensure we cover all dates
        data = yf.download(ticker, period="max", progress=False, auto_adjust=False)
        if data.empty:
            logger.error("Failed to fetch FX data. Returned empty dataset.")
            sys.exit(1)

        # Handle potential MultiIndex in columns
        if isinstance(data.columns, pd.MultiIndex):
            # If MultiIndex (Price, Ticker), we want (Close, Ticker) or just Close
            if ticker in data.columns.get_level_values(1):
                 fx_series = data['Close'][ticker]
            else:
                 # Fallback if ticker structure is different
                 fx_series = data['Close'].iloc[:, 0]
        else:
            fx_series = data['Close']

        # Ensure index is datetime and normalized to date (midnight)
        fx_series.index = pd.to_datetime(fx_series.index).normalize()

        # Remove timezone info to match the naive lookup dates in get_fx_rate
        if fx_series.index.tz is not None:
             fx_series.index = fx_series.index.tz_localize(None)

        # Sort index to ensure asof works
        fx_series = fx_series.sort_index()

        return fx_series
    except Exception as e:
        logger.error(f"Error fetching FX history: {e}")
        sys.exit(1)

def get_fx_rate(date_obj, fx_history):
    """
    Helper to look up rate for a given date.
    Uses 'asof' to find the most recent rate if exact match is missing (e.g. weekends).
    """
    # Normalize date to midnight and strip timezone to match fx_history index
    target_date = pd.Timestamp(date_obj).tz_localize(None).normalize()

    # asof returns the value at or before the label
    try:
        rate = fx_history.asof(target_date)
        if pd.isna(rate):
            # If date is before the start of history
            logger.warning(f"Date {target_date.date()} is before start of FX history ({fx_history.index.min().date()}). Using first available.")
            return fx_history.iloc[0]
        return rate
    except Exception as e:
        logger.warning(f"Error looking up FX rate for {target_date}: {e}. Defaulting to 1.0 (Safety).")
        return 1.0

def safe_decimal_conversion(val):
    """
    Safely converts a value to Decimal, handling commas and NaNs.
    """
    s = str(val).replace(',', '')
    try:
        return Decimal(s)
    except:
        return Decimal(0)

def normalize_columns(df):
    """
    Normalizes dataframe columns to uppercase for case-insensitive matching.
    Returns the dataframe with upper-case columns.
    """
    df.columns = [c.strip().upper() for c in df.columns]
    return df

def parse_coinbase(filepath):
    """
    Parses Coinbase CSV.
    """
    if not os.path.exists(filepath):
        logger.warning(f"Coinbase file not found at: {filepath}. Skipping.")
        return pd.DataFrame()

    logger.info(f"Processing Coinbase file: {filepath}")
    try:
        df = pd.read_csv(filepath)

        # Case insensitive column matching
        # Mapping requested:
        # 'Timestamp', 'Transaction Type', 'Asset', 'Quantity Transacted', 'Total (inclusive of fees and/or spread)'

        # Create a map of UPPERCASE -> Original/Standard names
        req_cols_map = {
            'TIMESTAMP': 'Timestamp',
            'TRANSACTION TYPE': 'Transaction Type',
            'ASSET': 'Asset',
            'QUANTITY TRANSACTED': 'Quantity Transacted',
            'TOTAL (INCLUSIVE OF FEES AND/OR SPREAD)': 'Total (inclusive of fees and/or spread)'
        }

        # Normalize df columns to upper to check existence
        df_upper_cols = {c.strip().upper(): c for c in df.columns}

        # Check for missing columns
        missing = []
        rename_map = {}

        for req_upper, req_std in req_cols_map.items():
            if req_upper not in df_upper_cols:
                missing.append(req_std)
            else:
                rename_map[df_upper_cols[req_upper]] = req_upper # Rename actual col to UPPER internal name temporarily

        if missing:
            logger.error(f"Coinbase file missing columns (case-insensitive check): {missing}")
            return pd.DataFrame()

        # Rename to the UPPER keys for consistent access
        df = df.rename(columns=rename_map)

        # Column Mapping (using the UPPER keys now)
        # 'TIMESTAMP' -> 'TIME_UTC'
        df['TIME_UTC'] = pd.to_datetime(df['TIMESTAMP'], utc=True)

        df = df.rename(columns={
            'TRANSACTION TYPE': 'TXN_TYPE_RAW',
            'ASSET': 'COIN',
            'QUANTITY TRANSACTED': 'QTY',
            'TOTAL (INCLUSIVE OF FEES AND/OR SPREAD)': 'TOTAL_USD'
        })

        df['EXCHANGE'] = 'Coinbase'

        # Filtering & Transaction Mapping - Case Insensitive
        # Normalize TXN_TYPE_RAW to UPPER
        df['TXN_TYPE_RAW'] = df['TXN_TYPE_RAW'].astype(str).str.strip().str.upper()

        # Map keys are all UPPER
        keep_map = {
            'ADVANCE TRADE BUY': 'BUY',
            'ADVANCED TRADE BUY': 'BUY',
            'BUY': 'BUY',
            'ADVANCE TRADE SELL': 'SELL',
            'ADVANCED TRADE SELL': 'SELL',
            'SELL': 'SELL',
            'REWARD INCOME': 'INCOME',
            'STAKING INCOME': 'INCOME'
        }

        df = df[df['TXN_TYPE_RAW'].isin(keep_map.keys())].copy()
        df['TXN_TYPE'] = df['TXN_TYPE_RAW'].map(keep_map)

        return df[['EXCHANGE', 'TIME_UTC', 'TXN_TYPE', 'COIN', 'QTY', 'TOTAL_USD']]

    except Exception as e:
        logger.error(f"Error parsing Coinbase file: {e}")
        return pd.DataFrame()

def parse_binance(filepath):
    """
    Parses Binance CSV.
    """
    if not os.path.exists(filepath):
        logger.warning(f"Binance file not found at: {filepath}. Skipping.")
        return pd.DataFrame()

    logger.info(f"Processing Binance file: {filepath}")
    try:
        df = pd.read_csv(filepath)

        # Case insensitive column matching
        # Req: 'Date(UTC)', 'Base Asset', 'Quote Asset', 'Type', 'Amount', 'Total'

        req_cols_map = {
            'DATE(UTC)': 'Date(UTC)',
            'BASE ASSET': 'Base Asset',
            'QUOTE ASSET': 'Quote Asset',
            'TYPE': 'Type',
            'AMOUNT': 'Amount',
            'TOTAL': 'Total'
        }

        df_upper_cols = {c.strip().upper(): c for c in df.columns}

        missing = []
        rename_map = {}

        for req_upper, req_std in req_cols_map.items():
            if req_upper not in df_upper_cols:
                # Try finding without (UTC) maybe? User said strict mapping but case insensitive.
                # Let's stick to the names provided but allow case variance.
                missing.append(req_std)
            else:
                rename_map[df_upper_cols[req_upper]] = req_upper

        if missing:
            logger.error(f"Binance file missing columns: {missing}")
            return pd.DataFrame()

        df = df.rename(columns=rename_map)

        df['TIME_UTC'] = pd.to_datetime(df['DATE(UTC)'], dayfirst=True, utc=True)

        df = df.rename(columns={
            'BASE ASSET': 'COIN',
            'QUOTE ASSET': 'CCY',
            'AMOUNT': 'QTY',
            'TOTAL': 'TOTAL_USD_RAW'
        })

        df['EXCHANGE'] = 'Binance'

        # Logic 1: KEEP only 'Buy', 'Sell' (Case insensitive)
        df['TYPE_NORM'] = df['TYPE'].astype(str).str.strip().str.upper()

        keep_map = {'BUY': 'BUY', 'SELL': 'SELL'}
        df = df[df['TYPE_NORM'].isin(keep_map.keys())].copy()
        df['TXN_TYPE'] = df['TYPE_NORM'].map(keep_map)

        # Logic 2: USDT Fix
        # "If CCY column is 'USDT', treat it as 'USD'."
        # Case insensitive check for CCY
        df['CCY_NORM'] = df['CCY'].astype(str).str.strip().str.upper()
        valid_ccy = ['USD', 'USDT']
        df = df[df['CCY_NORM'].isin(valid_ccy)].copy()

        df = df.rename(columns={'TOTAL_USD_RAW': 'TOTAL_USD'})

        return df[['EXCHANGE', 'TIME_UTC', 'TXN_TYPE', 'COIN', 'QTY', 'TOTAL_USD']]

    except Exception as e:
        logger.error(f"Error parsing Binance file: {e}")
        return pd.DataFrame()

def format_excel_file(filepath):
    """
    Applies custom formatting to the Excel file using openpyxl.
    - Remove gridlines.
    - Format QTY to 9 decimals.
    - Format Money columns to 2 decimals.
    - Autofit columns.
    """
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active

        # 1. Remove Gridlines
        ws.sheet_view.showGridLines = False

        # 2. Number Formatting
        # Define formats
        fmt_qty = '0.000000000'
        fmt_money = '#,##0.00'
        fmt_date = 'yyyy-mm-dd hh:mm:ss'

        # Identify columns by header
        # Headers are in row 1
        headers = {cell.value: cell.column for cell in ws[1]}

        # Columns to format
        qty_col = headers.get('QTY')
        money_cols = [headers.get(c) for c in ['TOTAL_USD', 'TOTAL_GBP', 'EFF_UNIT_PRICE_USD', 'EFF_UNIT_PRICE_GBP'] if c in headers]
        date_col = headers.get('TIME_UTC')

        # Iterate over rows (starting from row 2)
        for row in ws.iter_rows(min_row=2):
            # QTY
            if qty_col:
                row[qty_col-1].number_format = fmt_qty

            # Money
            for col_idx in money_cols:
                if col_idx and col_idx <= len(row): # check if col_idx is within bounds of row
                     # row is a tuple of cells. Index is 0-based. col_idx is 1-based.
                     row[col_idx-1].number_format = fmt_money

            # Date
            if date_col:
                row[date_col-1].number_format = fmt_date

        # 3. Autofit Columns
        for column_cells in ws.columns:
            # We need to calculate length, but be careful with None
            length = 0
            for cell in column_cells:
                val = cell.value
                if val is not None:
                    length = max(length, len(str(val)))

            # Add a little buffer
            adjusted_width = length + 2

            # Simple heuristic adjustment for formatted columns where raw string len might be small
            col_idx = column_cells[0].column

            if col_idx == qty_col:
                # 9 decimals + integer part. Safe min width ~15-20?
                if adjusted_width < 18:
                    adjusted_width = 18

            if col_idx in money_cols:
                if adjusted_width < 12:
                    adjusted_width = 12

            if col_idx == date_col:
                 if adjusted_width < 20:
                     adjusted_width = 20

            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

        wb.save(filepath)
        logger.info(f"Applied Excel formatting to {filepath}")

    except Exception as e:
        logger.error(f"Failed to apply Excel formatting: {e}")

def apply_signage_rules(row):
    """
    Applies sign enforcement logic.
    Returns modified QTY, TOTAL_USD, TOTAL_GBP (as Decimals).
    """
    txn_type = row['TXN_TYPE']

    # Ensure they are Decimals
    qty = row['QTY']
    total_usd = row['TOTAL_USD']
    total_gbp = row['TOTAL_GBP']

    # Absolute values first to reset
    qty = abs(qty)
    total_usd = abs(total_usd)
    total_gbp = abs(total_gbp)

    if txn_type in ['BUY', 'INCOME']:
        # QTY Positive (+)
        # TOTAL_USD and TOTAL_GBP Negative (-)
        qty = qty
        total_usd = -total_usd
        total_gbp = -total_gbp

    elif txn_type == 'SELL':
        # QTY Negative (-)
        # TOTAL_USD and TOTAL_GBP Positive (+)
        qty = -qty
        total_usd = total_usd
        total_gbp = total_gbp

    return qty, total_usd, total_gbp

def main():
    logger.info("Starting Master Crypto Data Sheet generation...")

    # 1. Fetch FX
    fx_history = fetch_fx_history()
    if fx_history.empty:
        logger.error("FX history is empty. Aborting.")
        return

    logger.info(f"FX History fetched. Range: {fx_history.index.min().date()} to {fx_history.index.max().date()}")

    # 2. Parse Data
    df_coinbase = parse_coinbase(COINBASE_FILE_PATH)
    df_binance = parse_binance(BINANCE_FILE_PATH)

    if df_coinbase.empty and df_binance.empty:
        logger.error("No data found from either source. Exiting.")
        return

    # 3. Merge
    master_df = pd.concat([df_coinbase, df_binance], ignore_index=True)
    logger.info(f"Combined data. Total rows: {len(master_df)}")

    # 4. Master Logic Loop
    processed_rows = []

    for idx, row in master_df.iterrows():
        # Clean numeric inputs first
        qty_dec = safe_decimal_conversion(row['QTY'])
        total_usd_dec = safe_decimal_conversion(row['TOTAL_USD'])

        # A. FX Rate
        rate_val = get_fx_rate(row['TIME_UTC'], fx_history)
        rate_dec = Decimal(str(rate_val))

        # B. GBP Conversion
        if rate_dec == 0:
            total_gbp_dec = Decimal(0)
        else:
            # TOTAL_GBP = TOTAL_USD / FX_RATE
            total_gbp_dec = total_usd_dec / rate_dec

        # Update row with preliminary values (as Decimals)
        row['QTY'] = qty_dec
        row['TOTAL_USD'] = total_usd_dec
        row['FX_RATE'] = rate_dec
        row['TOTAL_GBP'] = total_gbp_dec

        # C. Signage Enforcement
        qty, total_usd, total_gbp = apply_signage_rules(row)

        row['QTY'] = qty
        row['TOTAL_USD'] = total_usd
        row['TOTAL_GBP'] = total_gbp

        # D. Calculate EFF_UNIT_PRICE
        # EFF_UNIT_PRICE_USD = TOTAL_USD / QTY
        # EFF_UNIT_PRICE_GBP = TOTAL_GBP / QTY
        # Handle Division by Zero

        if qty == 0:
            row['EFF_UNIT_PRICE_USD'] = Decimal(0)
            row['EFF_UNIT_PRICE_GBP'] = Decimal(0)
        else:
            row['EFF_UNIT_PRICE_USD'] = total_usd / qty
            row['EFF_UNIT_PRICE_GBP'] = total_gbp / qty

        processed_rows.append(row)

    final_df = pd.DataFrame(processed_rows)

    # 5. Final Formatting
    cols = ['EXCHANGE', 'TIME_UTC', 'TXN_TYPE', 'COIN', 'QTY', 'TOTAL_USD', 'FX_RATE', 'TOTAL_GBP', 'EFF_UNIT_PRICE_USD', 'EFF_UNIT_PRICE_GBP']
    final_df = final_df[cols]

    # Sort
    final_df = final_df.sort_values(by='TIME_UTC', ascending=True)

    # Pre-Save: Remove timezone for Excel compatibility
    # Excel does not support timezone-aware datetimes.
    # We will keep the original final_df for CSV (which handles TZ strings fine),
    # but create a copy for Excel with TZ stripped (or converted to string).
    # Converting to string is safest to preserve the visual UTC offset info if needed,
    # but normally users just want the date/time.
    # Let's convert to TZ-naive (UTC) for Excel.

    final_df_xlsx = final_df.copy()
    if 'TIME_UTC' in final_df_xlsx.columns:
         final_df_xlsx['TIME_UTC'] = final_df_xlsx['TIME_UTC'].dt.tz_localize(None)

    # Save
    # Ensure directory exists
    output_dir = os.path.dirname(OUTPUT_FILE_NAME)
    if output_dir and not os.path.exists(output_dir):
        try:
            os.makedirs(output_dir)
        except OSError:
            pass # Might fail on permissions, but user config implies paths exist or are writable

    try:
        # Save CSV
        final_df.to_csv(OUTPUT_FILE_NAME, index=False)

        # Save XLSX
        # Replace .csv extension with .xlsx or append if no extension
        if OUTPUT_FILE_NAME.lower().endswith('.csv'):
            output_xlsx = OUTPUT_FILE_NAME[:-4] + '.xlsx'
        else:
            output_xlsx = OUTPUT_FILE_NAME + '.xlsx'

        final_df_xlsx.to_excel(output_xlsx, index=False)

        # Apply Excel Formatting
        format_excel_file(output_xlsx)

        # Summary
        n_coinbase = len(final_df[final_df['EXCHANGE'] == 'Coinbase'])
        n_binance = len(final_df[final_df['EXCHANGE'] == 'Binance'])
        print(f"\nSUCCESS. Total Rows: {len(final_df)}. Coinbase: {n_coinbase}, Binance: {n_binance}. Files saved to {OUTPUT_FILE_NAME} and {output_xlsx}.")
    except Exception as e:
        logger.error(f"Failed to save output file: {e}")

if __name__ == "__main__":
    main()
