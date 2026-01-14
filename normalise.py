import os
import sys
import json
import hashlib
import datetime
import shutil
import argparse
import csv
import io
import warnings
import re
import subprocess
from typing import List, Dict, Any, Optional, Tuple, Union

# Attempt imports for required libraries
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
    from docx.shared import Inches
except ImportError:
    docx = None

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    pptx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    from pyxlsb import open_workbook as open_xlsb
except ImportError:
    open_xlsb = None

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

# ---------------------------------------------------------------------------
# Globals & Config
# ---------------------------------------------------------------------------

CONFIG_FILENAME = "config.json"
REQUIRED_CONFIG_KEYS = [
    "input_root", "output_root", "subfolders", "images_folder_name",
    "log_filename", "overwrite", "max_table_rows"
]

TOKEN_LIMIT_CHUNK = 1000  # Target tokens per chunk

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def ensure_dependencies(config_path: str = CONFIG_FILENAME):
    if os.path.exists(config_path):
        base_dir = os.path.dirname(os.path.abspath(config_path))
        req_path = os.path.join(base_dir, "requirements.txt")
    else:
        req_path = "requirements.txt"

    if not os.path.exists(req_path):
        return

    missing = []
    if fitz is None: missing.append("pymupdf")
    if docx is None: missing.append("python-docx")
    if pptx is None: missing.append("python-pptx")
    if openpyxl is None: missing.append("openpyxl")
    if pd is None: missing.append("pandas")
    if open_xlsb is None: missing.append("pyxlsb")
    if Image is None: missing.append("pillow")
    if pytesseract is None: missing.append("pytesseract")

    if missing:
        print(f"Missing libraries: {', '.join(missing)}. Attempting to install from {req_path}...")
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "-r", req_path])
            print("Dependencies installed. Restarting script...")
            os.execv(sys.executable, [sys.executable] + sys.argv)
        except Exception as e:
            print(f"Failed to install dependencies: {e}")

def load_config(path: str = CONFIG_FILENAME) -> Dict[str, Any]:
    if not os.path.exists(path):
        return {
             "input_root": ".", "output_root": "output",
             "subfolders": [], "images_folder_name": "images",
             "log_filename": "processing_log.jsonl", "overwrite": True,
             "max_table_rows": 2000
        }
    try:
        with open(path, 'r', encoding='utf-8') as f:
            config = json.load(f)
        return config
    except json.JSONDecodeError as e:
        print(f"Error parsing config JSON: {e}")
        sys.exit(1)

def get_unique_id(source_relpath: str, location: str, object_index: int) -> str:
    norm_path = source_relpath.replace(os.sep, "/")
    raw = f"{norm_path}|{location}|{object_index}"
    return hashlib.sha256(raw.encode('utf-8')).hexdigest()[:16]

def clean_text(text: str) -> str:
    if not text: return ""
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def estimate_tokens(text: str) -> int:
    return len(text) // 4

def get_image_dimensions(image_bytes: bytes) -> Tuple[int, int]:
    if not Image: return (0, 0)
    try:
        with Image.open(io.BytesIO(image_bytes)) as img:
            return img.width, img.height
    except Exception:
        return (0, 0)

# ---------------------------------------------------------------------------
# Semantic Structures
# ---------------------------------------------------------------------------

class SemanticUnit:
    def __init__(self, obj_type, content, location, source, idx, **kwargs):
        self.obj_type = obj_type  # text, table, image
        self.content = content
        self.location = location
        self.source = source
        self.id = get_unique_id(source, location, idx)
        self.metadata = kwargs

    def to_dict(self):
        return {
            "object_id": self.id,
            "object_type": self.obj_type,
            "source_location": self.location,
            "content": self.content,
            **self.metadata
        }

def analyze_image(image_bytes: bytes, width: int, height: int, nearby_text: str, filename: str, config: Dict[str, Any]) -> Dict[str, Any]:
    # Classification
    cls_config = config.get("image_classification", {})
    min_w = cls_config.get("min_width", 100)
    min_h = cls_config.get("min_height", 100)
    keywords = cls_config.get("keywords", {})

    img_type = "conceptual_diagram"
    role = "illustration"
    confidence = 0.8

    # Dimensions check
    if width > 0 and height > 0:
        if width < min_w or height < min_h:
            ratio = width / height
            if ratio > 10 or ratio < 0.1:
                img_type = "decorative"
                role = "decorative"
                confidence = 0.9
            elif "logo" in filename.lower():
                img_type = "logo"
                role = "decorative"
            else:
                img_type = "decorative"
                role = "decorative"
                confidence = 0.6

    # Keywords
    text_lower = nearby_text.lower()
    caption_match = re.search(r'((?:figure|table|chart|graph|diagram|model)\s*\d+[:.]?\s*[^\n]+)', nearby_text, re.IGNORECASE)
    depicts = caption_match.group(1).strip() if caption_match else nearby_text.strip().replace('\n', ' ')
    if len(depicts) > 100: depicts = depicts[:100] + "..."
    if not depicts: depicts = "Visual content"

    if img_type == "conceptual_diagram":
        for k_type, k_words in keywords.items():
            if any(w in text_lower for w in k_words):
                img_type = k_type
                break
        if "chart" in img_type or "data" in img_type:
            role = "evidence"
        elif "formula" in img_type:
            role = "example"

    # OCR
    ocr_text = ""
    ocr_status = "not_attempted"
    if img_type not in ["decorative", "logo"]:
        if pytesseract and Image:
            try:
                img = Image.open(io.BytesIO(image_bytes))
                ocr_text = pytesseract.image_to_string(img).strip()
                ocr_status = "success" if ocr_text else "empty"
            except Exception:
                ocr_status = "failed"
        else:
            ocr_status = "lib_missing"

    return {
        "type": img_type,
        "depicts": depicts,
        "semantic_role": role,
        "confidence": confidence,
        "ocr_text": ocr_text,
        "ocr_status": ocr_status
    }

def save_image(image_data: bytes, ext: str, unique_id: str, config: Dict[str, Any]) -> str:
    images_dir = os.path.join(config["output_root"], config["images_folder_name"])
    os.makedirs(images_dir, exist_ok=True)
    if not ext: ext = "png"
    ext = ext.lstrip(".").lower()
    filename = f"IMG_{unique_id}.{ext}"
    filepath = os.path.join(images_dir, filename)
    with open(filepath, "wb") as f:
        f.write(image_data)
    return filename

def log_event(config: Dict[str, Any], entry: Dict[str, Any]):
    log_path = os.path.join(config["output_root"], config["log_filename"])
    os.makedirs(os.path.dirname(log_path), exist_ok=True)
    with open(log_path, 'a', encoding='utf-8') as f:
        f.write(json.dumps(entry) + "\n")

def get_metadata_header(source_path: str, source_relpath: str, source_type: str,
                        warnings_list: List[str], errors_list: List[str]) -> str:
    header = "--- METADATA ---\n"
    header += f"source_path: {source_path}\n"
    header += f"source_relpath: {source_relpath}\n"
    header += f"source_type: {source_type}\n"
    header += f"generated_utc: {datetime.datetime.now(datetime.timezone.utc).isoformat()}\n"

    toolchain = []
    if fitz: toolchain.append(f"PyMuPDF-{fitz.VersionBind}")
    if docx: toolchain.append("python-docx")
    if pptx: toolchain.append("python-pptx")
    if openpyxl: toolchain.append(f"openpyxl-{openpyxl.__version__}")
    if pd: toolchain.append(f"pandas-{pd.__version__}")
    if Image: toolchain.append(f"Pillow-{Image.__version__}")
    if pytesseract: toolchain.append(f"pytesseract")

    header += f"toolchain: {', '.join(toolchain)}\n"
    header += f"warnings: {json.dumps(warnings_list)}\n"
    header += f"errors: {json.dumps(errors_list)}\n"
    header += "----------------\n\n"
    return header

# ---------------------------------------------------------------------------
# Processors
# ---------------------------------------------------------------------------

def get_block_style(page, bbox: fitz.Rect) -> float:
    try:
        text_dict = page.get_text("dict", clip=bbox)
        max_size = 0
        for b in text_dict.get("blocks", []):
            for l in b.get("lines", []):
                for s in l.get("spans", []):
                    if s["size"] > max_size: max_size = s["size"]
        return max_size
    except:
        return 0

def process_pdf(file_path: str, rel_path: str, config: Dict[str, Any]) -> List[SemanticUnit]:
    if not fitz: raise ImportError("PyMuPDF missing")
    doc = fitz.open(file_path)
    units = []

    for page_idx, page in enumerate(doc):
        loc = f"PAGE-{page_idx+1}"
        page_units_with_pos = [] # (y, SemanticUnit)

        # 1. Extract Tables (High Priority)
        tables = []
        if hasattr(page, "find_tables"):
            try:
                tables_obj = page.find_tables()
                for i, tab in enumerate(tables_obj):
                    df = tab.to_pandas()
                    content = {
                        "columns": df.columns.tolist(),
                        "rows": df.values.tolist(),
                        "dtypes": [str(d) for d in df.dtypes],
                        "rows_original": len(df),
                        "rows_emitted": min(len(df), config["max_table_rows"]),
                        "truncated": len(df) > config["max_table_rows"]
                    }
                    if content["truncated"]:
                        content["rows"] = content["rows"][:config["max_table_rows"]]

                    bbox = fitz.Rect(tab.bbox)
                    unit = SemanticUnit("table", content, loc, rel_path, len(units) + len(page_units_with_pos), semantic_role="data")
                    tables.append({"rect": bbox, "unit": unit})
                    page_units_with_pos.append((bbox.y0, unit))
            except Exception as e:
                pass

        # 2. Extract Text (via blocks)
        # blocks: (x0, y0, x1, y1, text, block_no, block_type)
        blocks = page.get_text("blocks")

        for b in blocks:
            b_rect = fitz.Rect(b[0], b[1], b[2], b[3])

            # Skip if overlaps significantly with extracted table
            overlap = False
            for t in tables:
                intersection = b_rect & t["rect"]
                if intersection.get_area() > 0.9 * b_rect.get_area():
                    overlap = True; break
            if overlap: continue

            if b[6] == 0: # Text
                block_text = clean_text(b[4])
                if not block_text: continue

                # Check semantics via dict clip
                max_size = get_block_style(page, b_rect)

                role = "body"
                if max_size > 14: role = "heading"
                elif "?" in block_text and len(block_text) < 200: role = "question"

                unit = SemanticUnit("text", block_text, loc, rel_path, len(units) + len(page_units_with_pos),
                                    text_type=role, hierarchy_level=1 if role=="heading" else 2)
                page_units_with_pos.append((b_rect.y0, unit))

        # 3. Handle Images separately
        image_list = page.get_images(full=True)
        for img_idx, img in enumerate(image_list):
            xref = img[0]
            try:
                rects = page.get_image_rects(xref)
                for rect in rects:
                    # Check overlap with tables
                    overlap = False
                    for t in tables:
                        intersection = rect & t["rect"]
                        if intersection.get_area() > 0.9 * rect.get_area():
                            overlap = True; break
                    if overlap: continue

                    # Extract
                    base_image = doc.extract_image(xref)
                    image_bytes = base_image["image"]
                    ext = base_image["ext"]
                    w, h = get_image_dimensions(image_bytes)
                    if w == 0: w, h = int(rect.width), int(rect.height)

                    analysis = analyze_image(image_bytes, w, h, "", os.path.basename(file_path), config)
                    unique_id = get_unique_id(rel_path, loc, len(units) + len(page_units_with_pos))
                    save_image(image_bytes, ext, unique_id, config)

                    unit = SemanticUnit("image", {"filename": f"IMG_{unique_id}.{ext}"}, loc, rel_path, len(units) + len(page_units_with_pos), **analysis)
                    page_units_with_pos.append((rect.y0, unit))
            except: pass

        # Sort page units by vertical position
        page_units_with_pos.sort(key=lambda x: x[0])
        units.extend([u[1] for u in page_units_with_pos])

    return units

def process_docx(file_path: str, rel_path: str, config: Dict[str, Any]) -> List[SemanticUnit]:
    if not docx: raise ImportError("python-docx missing")
    doc = docx.Document(file_path)
    units = []

    def iter_elements(doc_obj):
        if isinstance(doc_obj, docx.document.Document): body = doc_obj.element.body
        else: return
        for child in body.iterchildren():
            if child.tag.endswith('p'): yield docx.text.paragraph.Paragraph(child, doc_obj)
            elif child.tag.endswith('tbl'): yield docx.table.Table(child, doc_obj)

    loc = "SECTION-1"

    for idx, block in enumerate(iter_elements(doc)):
        if isinstance(block, docx.text.paragraph.Paragraph):
            text = clean_text(block.text)
            if not text and not block.runs: continue # Skip empty unless images?

            # Images
            for run in block.runs:
                blips = run.element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                for blip in blips:
                    embed = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if embed:
                        try:
                            part = doc.part.related_parts[embed]
                            img_bytes = part.blob
                            ct = part.content_type
                            ext = ct.split('/')[-1] if '/' in ct else 'png'
                            w, h = get_image_dimensions(img_bytes)

                            analysis = analyze_image(img_bytes, w, h, text, os.path.basename(file_path), config)
                            unique_id = get_unique_id(rel_path, loc, len(units))
                            save_image(img_bytes, ext, unique_id, config)
                            units.append(SemanticUnit("image", {"filename": f"IMG_{unique_id}.{ext}"}, loc, rel_path, len(units), **analysis))
                        except: pass

            if text:
                role = "body"
                style = block.style.name.lower()
                if "heading" in style or "title" in style: role = "heading"
                elif "list" in style: role = "list_item"
                elif "?" in text and len(text) < 150: role = "question"

                units.append(SemanticUnit("text", text, loc, rel_path, len(units), text_type=role, hierarchy_level=1 if role=="heading" else 2))

        elif isinstance(block, docx.table.Table):
            rows = []
            for row in block.rows:
                rows.append([cell.text.strip() for cell in row.cells])

            if not rows: continue

            header = rows[0]
            data = rows[1:] if len(rows) > 1 else []

            content = {
                "columns": header,
                "rows": data,
                "dtypes": ["string"] * len(header),
                "rows_original": len(rows),
                "rows_emitted": min(len(rows), config["max_table_rows"]),
                "truncated": len(rows) > config["max_table_rows"]
            }
            if content["truncated"]:
                content["rows"] = content["rows"][:config["max_table_rows"]-1]

            units.append(SemanticUnit("table", content, loc, rel_path, len(units), semantic_role="data"))

    return units

def process_pptx(file_path: str, rel_path: str, config: Dict[str, Any]) -> List[SemanticUnit]:
    if not pptx: raise ImportError("python-pptx missing")
    prs = pptx.Presentation(file_path)
    units = []

    for slide_idx, slide in enumerate(prs.slides):
        loc = f"SLIDE-{slide_idx+1}"
        shapes = sorted(slide.shapes, key=lambda s: (s.top, s.left))

        slide_text = " ".join([s.text for s in shapes if hasattr(s, "text") and s.text])

        for shape in shapes:
            if hasattr(shape, "text") and shape.text:
                text = clean_text(shape.text)
                if text:
                    role = "body"
                    # Try to detect title from placeholder
                    if shape.is_placeholder and "TITLE" in str(shape.placeholder_format.type):
                        role = "heading"
                    units.append(SemanticUnit("text", text, loc, rel_path, len(units), text_type=role, hierarchy_level=1 if role=="heading" else 2))

            if shape.has_table:
                rows = []
                for row in shape.table.rows:
                    rows.append([cell.text_frame.text.strip() for cell in row.cells])
                if rows:
                    content = {
                        "columns": rows[0],
                        "rows": rows[1:],
                        "dtypes": ["string"] * len(rows[0]),
                        "rows_original": len(rows),
                        "rows_emitted": min(len(rows), config["max_table_rows"]),
                        "truncated": len(rows) > config["max_table_rows"]
                    }
                    if content["truncated"]: content["rows"] = content["rows"][:config["max_table_rows"]-1]
                    units.append(SemanticUnit("table", content, loc, rel_path, len(units), semantic_role="data"))

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img_bytes = shape.image.blob
                    ext = shape.image.ext
                    w = int(shape.width / 9525)
                    h = int(shape.height / 9525)

                    analysis = analyze_image(img_bytes, w, h, slide_text, os.path.basename(file_path), config)
                    unique_id = get_unique_id(rel_path, loc, len(units))
                    save_image(img_bytes, ext, unique_id, config)
                    units.append(SemanticUnit("image", {"filename": f"IMG_{unique_id}.{ext}"}, loc, rel_path, len(units), **analysis))
                except: pass

    return units

def process_excel(file_path: str, rel_path: str, config: Dict[str, Any]) -> List[SemanticUnit]:
    units = []
    ext = os.path.splitext(file_path)[1].lower()

    def extract_rows(sheet_name, rows_iter):
        rows = list(rows_iter)
        if not rows: return
        header = [str(c) if c is not None else "" for c in rows[0]]
        data = [[str(c) if c is not None else "" for c in r] for r in rows[1:]]

        content = {
            "columns": header,
            "rows": data,
            "dtypes": ["string"] * len(header), # Inference is hard without pandas
            "rows_original": len(rows),
            "rows_emitted": min(len(rows), config["max_table_rows"]),
            "truncated": len(rows) > config["max_table_rows"]
        }
        if content["truncated"]: content["rows"] = content["rows"][:config["max_table_rows"]-1]

        units.append(SemanticUnit("table", content, f'SHEET "{sheet_name}"', rel_path, len(units), semantic_role="data"))

    if ext == '.csv':
        try:
            with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
                extract_rows("CSV", csv.reader(f))
        except: pass
    elif ext == '.xlsx' and openpyxl:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                extract_rows(sheet, [[c.value for c in r] for r in ws.rows])
        except: pass
    elif ext == '.xlsb' and open_xlsb:
        try:
            with open_xlsb(file_path) as wb:
                for sheet in wb.sheets:
                    with wb.get_sheet(sheet) as ws:
                        extract_rows(sheet, [[c.v for c in r] for r in ws.rows()])
        except: pass

    return units

# ---------------------------------------------------------------------------
# Core Pipeline
# ---------------------------------------------------------------------------

def render_output(units: List[SemanticUnit], output_path: str):
    # 1. Concept Index
    concepts = set()
    for u in units:
        if u.obj_type == "text" and u.metadata.get("text_type") == "heading":
            concepts.add(u.content.strip())
        if u.obj_type == "image":
            depicts = u.metadata.get("depicts", "")
            if len(depicts) > 5 and "visual content" not in depicts.lower():
                concepts.add(depicts[:50])

    # 2. Chunking
    chunks = []
    current_chunk = {"id": 0, "tokens": 0, "items": []}

    for u in units:
        # Estimate tokens
        cost = 0
        if u.obj_type == "text": cost = estimate_tokens(u.content)
        elif u.obj_type == "table": cost = 200 # rough
        elif u.obj_type == "image": cost = 100

        # Check limit (soft)
        if current_chunk["tokens"] + cost > TOKEN_LIMIT_CHUNK and current_chunk["items"]:
            chunks.append(current_chunk)
            current_chunk = {"id": len(chunks), "tokens": 0, "items": []}

        current_chunk["items"].append(u)
        current_chunk["tokens"] += cost

    if current_chunk["items"]: chunks.append(current_chunk)

    # 3. Write
    with open(output_path, 'w', encoding='utf-8') as f:
        # File Metadata
        f.write("--- FILE METADATA ---\n")
        f.write(json.dumps({
            "generated_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "concepts": list(concepts)[:10],
            "chunk_count": len(chunks)
        }, indent=2) + "\n\n")

        for ch in chunks:
            f.write(f"--- CHUNK {ch['id']} (est. {ch['tokens']} tokens) ---\n")
            f.write(json.dumps({
                "chunk_id": str(ch["id"]),
                "estimated_tokens": ch["tokens"],
                "modalities": list(set(i.obj_type for i in ch["items"])),
                "importance": "high" if any(i.metadata.get("text_type")=="heading" for i in ch["items"]) else "medium"
            }) + "\n")

            for item in ch["items"]:
                # Markdown Render
                if item.obj_type == "text":
                    prefix = "#" * item.metadata.get("hierarchy_level", 1) + " " if item.metadata.get("text_type") == "heading" else ""
                    f.write(f"\n{prefix}{item.content}\n")
                elif item.obj_type == "table":
                    f.write(f"\n[TABLE: {item.source} {item.location}]\n")
                    # Render Markdown Table
                    cols = item.content["columns"]
                    f.write("| " + " | ".join(map(str, cols)) + " |\n")
                    f.write("| " + " | ".join(["---"]*len(cols)) + " |\n")
                    for row in item.content["rows"]:
                        f.write("| " + " | ".join(map(str, row)) + " |\n")
                    if item.content.get("truncated"):
                        f.write(f"[TABLE_TRUNCATED: {item.content['rows_original']} rows]\n")
                elif item.obj_type == "image":
                    f.write(f"\n![{item.metadata.get('depicts')}]({item.content['filename']})\n")
                    f.write(f"[Image Role: {item.metadata.get('semantic_role')}]\n")

                # JSON Structure Block
                f.write("\n```json\n")
                f.write(json.dumps(item.to_dict(), indent=2))
                f.write("\n```\n")

            f.write("\n")

def process_file(source_abs: str, input_root: str, output_root: str, config: Dict[str, Any]) -> Dict[str, Any]:
    start = datetime.datetime.now()
    rel_path = os.path.relpath(source_abs, input_root)
    base, ext = os.path.splitext(rel_path)
    out_path = os.path.join(output_root, base + ".txt")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)

    file_ext = os.path.splitext(source_abs)[1].lower()
    units = []
    error = None

    try:
        if file_ext == ".pdf": units = process_pdf(source_abs, rel_path, config)
        elif file_ext in [".docx"]: units = process_docx(source_abs, rel_path, config)
        elif file_ext in [".pptx"]: units = process_pptx(source_abs, rel_path, config)
        elif file_ext in [".xlsx", ".csv", ".xlsb"]: units = process_excel(source_abs, rel_path, config)
        else: error = "Unsupported format"
    except Exception as e:
        error = str(e)

    status = "failed"
    if units:
        render_output(units, out_path)
        status = "success"
    elif error and "Unsupported" in error:
        status = "unsupported"
    elif error:
        status = "failed"

    # Quality Flags
    flags = {
        "table_truncated": any(u.obj_type=="table" and u.content.get("truncated") for u in units),
        "image_ocr_missing": any(u.obj_type=="image" and u.metadata.get("ocr_status")=="failed" for u in units),
        "low_image_confidence": any(u.obj_type=="image" and u.metadata.get("confidence", 1) < 0.5 for u in units)
    }

    if status == "success" and any(flags.values()):
        status = "success_partial"

    log_entry = {
        "source_path": source_abs,
        "status": status,
        "output_path": out_path,
        "units_count": len(units),
        "error": error,
        "quality_flags": flags,
        "duration_ms": int((datetime.datetime.now() - start).total_seconds() * 1000)
    }
    log_event(config, log_entry)
    return log_entry

def scan_and_process(config: Dict[str, Any]):
    print(f"Scanning {config['input_root']}...")
    stats = {"success": 0, "partial": 0, "failed": 0, "unsupported": 0}
    for folder in config["subfolders"]:
        path = os.path.join(config["input_root"], folder)
        if not os.path.exists(path): continue
        for root, _, files in os.walk(path):
            for f in files:
                fpath = os.path.join(root, f)
                print(f"Processing {f}...")
                res = process_file(fpath, config["input_root"], config["output_root"], config)

                if res["status"] == "success":
                    stats["success"] += 1
                elif res["status"] == "success_partial":
                    stats["partial"] += 1
                    print(f"  [PARTIAL] Quality issues: {json.dumps(res.get('quality_flags'))}")
                elif res["status"] == "unsupported":
                    stats["unsupported"] += 1
                    print(f"  [UNSUPPORTED] {res.get('error') or 'Type not supported'}")
                else:
                    stats["failed"] += 1
                    print(f"  [FAILED] {res.get('error')}")

    print(f"Done. Success: {stats['success']}, Partial: {stats['partial']}, Failed: {stats['failed']}, Unsupported: {stats['unsupported']}")

def self_test(config: Dict[str, Any]):
    print("Running self-test...")
    # 1. Classification Logic
    print("Testing Classification...")
    res = analyze_image(b"fake", 500, 500, "Figure 1: Sales Growth Chart", "img.png", config)
    assert res["type"] == "data_chart", f"Expected data_chart, got {res['type']}"
    assert "Sales Growth" in res["depicts"], f"Expected description capture, got {res['depicts']}"

    res = analyze_image(b"fake", 10, 10, "", "icon.png", config)
    assert res["type"] == "decorative", f"Expected decorative, got {res['type']}"

    # 2. Config
    assert config["image_classification"]["min_width"] == 100
    print("Config & Logic Check Passed.")

    # 3. Libraries
    print(f"Pillow available: {Image is not None}")

    print("Self-test completed successfully.")

if __name__ == "__main__":
    ensure_dependencies()

    parser = argparse.ArgumentParser()
    parser.add_argument("--self-test", action="store_true", help="Run self-test mode")
    args = parser.parse_args()

    cfg = load_config()

    if args.self_test:
        self_test(cfg)
    else:
        scan_and_process(cfg)
