import os
import openpyxl
from openpyxl.styles import Font

class LayoutConfig:
    def __init__(self, filepath):
        self.filepath = filepath
        self.tiles = []
        self.settings = {
            "missing_value": "BLANK", # BLANK or 0
            "output_structure": "PER_TILE" # PER_TILE or PER_SOURCE
        }

    def exists(self):
        return os.path.exists(self.filepath)

    def generate_default(self, dataset_loader=None):
        wb = openpyxl.Workbook()

        # 1. TILES Sheet
        ws_tiles = wb.active
        ws_tiles.title = "TILES"
        headers = [
            "Enabled (Y/N)", "Tile Order", "Title", "File Alias",
            "Sheet Name", "Row Label", "Column Labels", "Layout Mode",
            "Output Mode", "Output Sheet Name", "Notes"
        ]
        ws_tiles.append(headers)
        # Style headers
        for cell in ws_tiles[1]:
            cell.font = Font(bold=True)

        # 2. SETTINGS Sheet
        ws_settings = wb.create_sheet("SETTINGS")
        ws_settings.append(["Setting", "Value", "Description"])
        ws_settings.append(["missing_value", "BLANK", "Cell content when status != ok (BLANK or 0)"])
        ws_settings.append(["output_structure", "PER_TILE", "Grouping: PER_TILE (default) or PER_SOURCE"])
        for cell in ws_settings[1]:
            cell.font = Font(bold=True)

        # 3. CATALOGS
        ws_sources = wb.create_sheet("CATALOG_SOURCES")
        ws_sources.append(["File Alias", "Sheet Name"])
        for cell in ws_sources[1]:
            cell.font = Font(bold=True)

        ws_labels = wb.create_sheet("CATALOG_LABELS")
        ws_labels.append(["Row Label", "Column Label"])
        for cell in ws_labels[1]:
            cell.font = Font(bold=True)

        if dataset_loader:
            # Populate catalogs
            for (fa, sn) in sorted(list(dataset_loader.summary["sources"])):
                ws_sources.append([fa, sn])
            for (rl, cl) in sorted(list(dataset_loader.summary["labels"])):
                ws_labels.append([rl, cl])

        wb.save(self.filepath)
        print(f"Generated default layout at {self.filepath}")

    def load(self):
        if not self.exists():
            raise FileNotFoundError(f"Layout file not found: {self.filepath}")

        wb = openpyxl.load_workbook(self.filepath, data_only=True)

        # Load SETTINGS
        if "SETTINGS" in wb.sheetnames:
            ws = wb["SETTINGS"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None: continue
                key = str(row[0]).strip().lower()
                val = str(row[1]).strip() if row[1] is not None else ""

                if key == "missing_value":
                    if val.upper() in ["BLANK", "0"]:
                        self.settings["missing_value"] = val.upper()
                elif key == "output_structure":
                    if val.upper() in ["PER_TILE", "PER_SOURCE"]:
                        self.settings["output_structure"] = val.upper()

        # Load TILES
        if "TILES" not in wb.sheetnames:
            raise ValueError("Layout file must contain a 'TILES' sheet.")

        ws = wb["TILES"]
        headers = [c.value for c in ws[1]]

        # Map headers to indices
        col_map = {}
        expected_cols = {
            "Enabled (Y/N)": "enabled",
            "Tile Order": "order",
            "Title": "title",
            "File Alias": "file_alias",
            "Sheet Name": "sheet_name",
            "Row Label": "row_label",
            "Column Labels": "column_labels",
            "Layout Mode": "layout_mode",
            "Output Mode": "output_mode",
            "Output Sheet Name": "output_sheet_name"
        }

        for idx, h in enumerate(headers):
            if h in expected_cols:
                col_map[expected_cols[h]] = idx

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row: continue

            # Check enabled
            enabled_idx = col_map.get("enabled")
            if enabled_idx is not None:
                val = row[enabled_idx]
                if str(val).strip().upper() != "Y":
                    continue
            else:
                # If column missing, assume enabled? Or fail?
                # Prompt says "Enabled (Y/N)" is a column. If missing, we probably shouldn't crash, but assume Y.
                pass

            tile = {}
            # Extract fields
            for field, idx in col_map.items():
                tile[field] = row[idx]

            # Validation / Cleanup
            if not tile.get("row_label"):
                print("Warning: Skipping tile with missing Row Label.")
                continue

            # Defaults
            if not tile.get("layout_mode"): tile["layout_mode"] = "COLUMNS"
            if not tile.get("output_mode"): tile["output_mode"] = "VALUES"
            if not tile.get("output_sheet_name"): tile["output_sheet_name"] = "MI_Output"

            # Handle tile order (convert to int or float for sorting)
            try:
                tile["order"] = float(tile["order"]) if tile.get("order") is not None else 9999
            except:
                tile["order"] = 9999

            self.tiles.append(tile)

        # Sort by order
        self.tiles.sort(key=lambda x: x["order"])
