import os
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill
import re

class ExcelRenderer:
    def __init__(self, config):
        self.config = config
        self.missing_val = self.config.settings["missing_value"] # "BLANK" or "0"
        self.file_check_cache = {}

    def render(self, grouped_data, output_path):
        wb = openpyxl.Workbook()
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        # If no data, create a placeholder
        if not grouped_data:
            wb.create_sheet("Empty_Output")

        for sheet_name, tiles in grouped_data.items():
            # Clean sheet name for Excel (max 31 chars, no invalid chars)
            safe_sheet_name = self._sanitize_sheet_name(sheet_name)

            # Check for name collision in workbook
            base_name = safe_sheet_name
            count = 1
            while safe_sheet_name in wb.sheetnames:
                safe_sheet_name = f"{base_name[:28]}_{count}"
                count += 1

            ws = wb.create_sheet(safe_sheet_name)
            current_row = 1

            for tile in tiles:
                layout_mode = tile["config"]["layout_mode"].upper()

                if layout_mode == "SHEETS":
                    # Special case: Spawn new sheets
                    self._render_sheets_mode(wb, tile)
                else:
                    # Standard rendering
                    rows_used = self._render_tile(ws, current_row, tile)
                    current_row += rows_used + 2 # Add spacing

        wb.save(output_path)

    def _render_tile(self, ws, start_row, tile):
        """Renders a single tile into the given worksheet."""
        conf = tile["config"]
        layout_mode = conf["layout_mode"].upper()
        output_mode = conf["output_mode"].upper()

        # Title
        ws.cell(row=start_row, column=1, value=conf["title"]).font = Font(bold=True, size=12)
        start_row += 1

        columns = tile["columns"]
        data = tile["data"] # List of records matching columns
        row_label = conf["row_label"]

        if layout_mode == "COLUMNS":
            # Header Row: [Row Label] [Col 1] [Col 2] ...
            ws.cell(row=start_row, column=1, value="Row Label").font = Font(bold=True)
            for idx, col_name in enumerate(columns):
                cell = ws.cell(row=start_row, column=idx+2, value=col_name)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')

            start_row += 1

            # Data Row: [Label] [Val 1] [Val 2] ...
            ws.cell(row=start_row, column=1, value=row_label)
            for idx, record in enumerate(data):
                val = self._get_value(record, output_mode)
                ws.cell(row=start_row, column=idx+2, value=val)

            return 2 # Used 2 rows (Header + Data)

        elif layout_mode == "ROWS":
            # Header: [Column Label] [Value]
            ws.cell(row=start_row, column=1, value="Metric / Date").font = Font(bold=True)
            ws.cell(row=start_row, column=2, value=row_label).font = Font(bold=True)
            start_row += 1

            # Data Rows
            for idx, col_name in enumerate(columns):
                record = data[idx]
                val = self._get_value(record, output_mode)
                ws.cell(row=start_row+idx, column=1, value=col_name)
                ws.cell(row=start_row+idx, column=2, value=val)

            return len(columns) + 1

        return 0

    def _render_sheets_mode(self, wb, tile):
        """
        Spawns a new sheet for each column label.
        Naming: {OutputSheetName}_{ColumnLabel} or {File}-{Sheet}_{ColumnLabel}
        """
        conf = tile["config"]
        output_mode = conf["output_mode"].upper()
        base_name = conf["output_sheet_name"]
        if not base_name or base_name == "MI_Output":
             base_name = f"{conf['file_alias']}-{conf['sheet_name']}"

        columns = tile["columns"]
        data = tile["data"]
        row_label = conf["row_label"]

        for idx, col_name in enumerate(columns):
            sheet_name_raw = f"{base_name}_{col_name}"
            safe_name = self._sanitize_sheet_name(sheet_name_raw)

            # Handle collision
            count = 1
            orig_safe = safe_name
            while safe_name in wb.sheetnames:
                safe_name = f"{orig_safe[:28]}_{count}"
                count += 1

            ws = wb.create_sheet(safe_name)

            # Write data (Just a simple table for that single point?)
            # Or should it be a full tile-like structure?
            # Prompt doesn't specify layout *inside* the sheet for SHEETS mode.
            # Assuming standard "Header + Value".

            ws.cell(row=1, column=1, value=conf["title"]).font = Font(bold=True, size=12)

            ws.cell(row=3, column=1, value="Row Label").font = Font(bold=True)
            ws.cell(row=3, column=2, value=col_name).font = Font(bold=True)

            ws.cell(row=4, column=1, value=row_label)
            val = self._get_value(data[idx], output_mode)
            ws.cell(row=4, column=2, value=val)


    def _get_value(self, record, output_mode):
        if not record:
            return 0 if self.missing_val == "0" else None

        status = record.get("status", "ok")
        if status != "ok":
             return 0 if self.missing_val == "0" else None

        if output_mode == "LINKS":
            return self._create_link_formula(record)
        else:
            return record.get("value")

    def _create_link_formula(self, record):
        # Format: ='C:\Path\[File.xlsx]Sheet'!Cell
        path = record.get("link_file_path") or record.get("file_path")
        sheet = record.get("link_sheet_name") or record.get("sheet_name")
        cell = record.get("link_cell_address") or record.get("intersection_cell_address")

        # Prefer merged anchor if available
        if record.get("is_merged") and record.get("merged_anchor_address"):
            cell = record.get("merged_anchor_address")

        if not path or not sheet or not cell:
            return record.get("value") # Fallback

        # Check file existence (cached)
        if path not in self.file_check_cache:
            self.file_check_cache[path] = os.path.exists(path)
            if not self.file_check_cache[path]:
                print(f"Warning: Linked file not found: {path}")

        # Split path into dir and filename for Excel formula
        directory, filename = os.path.split(path)

        # Escape single quotes in sheet name by doubling them
        sheet = sheet.replace("'", "''")

        # Excel External Reference Syntax
        # If path has spaces, it technically doesn't need quotes in the formula IF using the bracket syntax?
        # Actually standard is: ='C:\Path\[Workbook.xlsx]SheetName'!A1
        # If spaces, wrapped in single quotes: ='C:\Path to stuff\[Workbook.xlsx]Sheet Name'!A1

        full_path = f"{directory}\\[{filename}]{sheet}"
        formula = f"='{full_path}'!{cell}"
        return formula

    def _sanitize_sheet_name(self, name):
        # Replace invalid chars: \ / ? * [ ] :
        invalid = r'[\\/\?\*\[\]\:]'
        name = re.sub(invalid, '_', str(name))
        return name[:31] # Max 31 chars
